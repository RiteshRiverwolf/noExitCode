"""A header field the reader could not find must never crash a note or print as "None".

    .venv\\Scripts\\python bench\\stage2\\missing_header_test.py

Found 2026-09-13 while regressing a fix to heading matching: on insp_1005 and
insp_1008 the Inspection Type was not read, and the code summary crashed on
`None.lower()`. The graph turned the crash into a stop for a person -- safe, but
over a field no decision depends on, and OCR can legitimately miss a label on a
real scan. Reading the rest of the code then showed the same gap elsewhere:
the model's instructions would have said "Inspection: None on None", and the
note would have printed "None barg" -- or crashed on `report_no.split()`.

That last crash was found by reading, not by watching it fail; this test is
what now demonstrates the code survives it.

Every header field is blanked at once, then:
  1. the code-written summary is produced, with no "None" in it;
  2. the model's instructions are produced, with no "None" in them;
  3. the Word note is rendered, contains no "None", and passes its read-back.
The decision itself does not depend on header fields, so it must not change.
"""

from __future__ import annotations

import copy
import re
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402

from workbench import prose, report_writer, rules  # noqa: E402
from workbench.evidence import HEADER_FIELDS, from_ground_truth  # noqa: E402

NONE_WORD = re.compile(r"\bNone\b")
failures: list[str] = []


def expect(name: str, ok: bool, detail: str) -> None:
    print(f"  {'pass' if ok else 'FAIL'}  {name}: {detail}")
    if not ok:
        failures.append(name)


full = from_ground_truth("insp_1002")
blank = copy.deepcopy(full)
for key in HEADER_FIELDS:
    blank.header[key] = None

decision = rules.evaluate(blank)
expect("the decision does not depend on header fields",
       decision.outcome == rules.evaluate(full).outcome, decision.outcome)

try:
    text = prose.code_summary(blank, decision)
    expect("code summary with every header field missing", not NONE_WORD.search(text), text[:110] + "...")
except Exception as e:
    expect("code summary with every header field missing", False, f"crashed: {type(e).__name__}: {e}")

try:
    facts = prose.facts_for_model(blank, decision)
    hits = [line for line in facts.splitlines() if NONE_WORD.search(line)]
    expect("the model's instructions", not hits, hits[0] if hits else "no 'None' in them")
except Exception as e:
    expect("the model's instructions", False, f"crashed: {type(e).__name__}: {e}")

with tempfile.TemporaryDirectory() as tmp:
    note = Path(tmp) / "note.docx"
    run_info = {"run_id": "missing-header-test", "generated_at": "13 Sep 2026 00:00",
                "graph": "inspection (test)", "trace": []}
    summary = prose.Summary(prose.code_summary(blank, decision), "code (test)")
    try:
        report_writer.render(blank, decision, summary, run_info, note)
        doc = Document(note)
        words = "\n".join([p.text for p in doc.paragraphs]
                          + [c.text for t in doc.tables for row in t.rows for c in row.cells]
                          + [p.text for p in doc.sections[0].header.paragraphs])
        hits = sorted({line.strip() for line in words.splitlines() if NONE_WORD.search(line)})
        expect("the rendered note", not hits, hits[0] if hits else "no 'None' anywhere in the file")
        results = report_writer.verify(note, blank, decision)
        failed = [n for n, ok, _ in results if not ok]
        expect("the note's read-back checks", not failed,
               f"{len(results) - len(failed)}/{len(results)} passed" + (f"; failed: {failed[0]}" if failed else ""))
    except Exception as e:
        expect("the rendered note", False, f"crashed: {type(e).__name__}: {e}")

print(f"\n{'all checks passed' if not failures else f'{len(failures)} check(s) FAILED'}")
sys.exit(1 if failures else 0)
