"""Stage 0 integration test: drive AnythingLLM entirely through its developer API.

Steps, in order:
  1. The API key is accepted.
  2. A workspace is created and pointed at the local agent model.
  3. A scanned inspection report is uploaded and embedded -- and we record
     whether AnythingLLM extracted any text from the scan at all.
  4. Test A (plumbing): an agent run started through the API calls our MCP
     tool and a Word file lands on the host. Almost no reasoning is needed, so
     a failure here is an integration failure, not a model failure.
  5. Test B (realistic): the agent must find the report in the workspace and
     write an approval note with the same tool. Passing requires the note's
     CONTENT to be right: it must name the below-minimum reading from the
     report's ground truth, with no template placeholders. A file merely
     existing is not a pass -- run 1 showed why.
  6. Test C (document Q&A, no agent): the same trap question through plain
     retrieval. This is the path AnythingLLM owns in our architecture.

The Word-file check relies on our own tool log (results/stage0/tool_calls.jsonl)
and the files on disk, not on what the agent says it did.

Environment:
  STAGE0_LABEL           tag for the result file (default "online")
  STAGE0_MODE            "upload-only" stops after the scan upload
  STAGE0_UPLOAD_TIMEOUT  seconds to wait for an upload (default 600)

    .venv\\Scripts\\python stage0\\run_integration_test.py

Writes results/stage0/run_<label>_<timestamp>.json.
"""

from __future__ import annotations

import json
import os
import re
import time
from datetime import datetime, timezone
from pathlib import Path

import requests
from docx import Document

ROOT = Path(__file__).resolve().parent.parent
BASE = os.environ.get("ALLM_BASE", "http://127.0.0.1:3001/api")
KEY_FILE = ROOT / "stage0" / "anythingllm" / "api_key.txt"
TOOL_LOG = ROOT / "results" / "stage0" / "tool_calls.jsonl"
OUT_DIR = ROOT / "stage0" / "output"
RESULTS = ROOT / "results" / "stage0"

# The trap case: every finding is Minor/Observation, one CML is below minimum.
SCAN = ROOT / "data" / "corpus" / "scans" / "medium" / "insp_1002_p1.png"
PDF = ROOT / "data" / "corpus" / "pdf" / "insp_1002.pdf"
TRUTH = ROOT / "data" / "corpus" / "truth" / "insp_1002.json"
MODEL = "llama3.1:8b"
CHAT_TIMEOUT = 900
UPLOAD_TIMEOUT = float(os.environ.get("STAGE0_UPLOAD_TIMEOUT", "600"))
PLACEHOLDER = re.compile(r"\[(list|insert|add|enter|todo)[^\]]*\]", re.IGNORECASE)


def clip(obj, n: int = 2000):
    """Truncate long strings anywhere in a JSON-like structure."""
    if isinstance(obj, str):
        return obj if len(obj) <= n else obj[:n] + f"... [+{len(obj) - n} chars]"
    if isinstance(obj, list):
        return [clip(x, n) for x in obj]
    if isinstance(obj, dict):
        return {k: clip(v, n) for k, v in obj.items()}
    return obj


def expected_facts() -> dict:
    """What a correct note or answer must contain, from the corpus ground truth."""
    t = json.loads(TRUTH.read_text(encoding="utf-8"))
    breached = [r for r in t["readings"] if r["is_below_minimum"]]
    return {
        "tag": t["equipment_tag"],
        "breached": [
            (r["cml_id"], f"{r['current_mm']:.2f}", f"{r['min_required_mm']:.1f}")
            for r in breached
        ],
    }


def content_checks(text: str, facts: dict) -> dict:
    checks = {
        f"mentions {facts['tag']}": facts["tag"] in text,
        "no template placeholders": PLACEHOLDER.search(text) is None,
    }
    for cml, current, minimum in facts["breached"]:
        checks[f"names breached {cml}"] = cml in text
        checks[f"gives {cml} current {current} mm"] = current in text
        checks[f"gives {cml} minimum {minimum} mm"] = minimum in text
    return checks


class Api:
    def __init__(self, key: str):
        self.session = requests.Session()
        self.session.headers["Authorization"] = f"Bearer {key}"

    def call(self, method: str, path: str, timeout: float = 120, **kw):
        t0 = time.perf_counter()
        try:
            r = self.session.request(method, BASE + path, timeout=timeout, **kw)
        except requests.RequestException as e:
            return None, f"{type(e).__name__}: {e}", round(time.perf_counter() - t0, 2)
        seconds = round(time.perf_counter() - t0, 2)
        try:
            body = r.json()
        except ValueError:
            body = r.text
        return r.status_code, body, seconds


def tool_log_len() -> int:
    if not TOOL_LOG.exists():
        return 0
    return len(TOOL_LOG.read_text(encoding="utf-8").splitlines())


def tool_calls_since(offset: int) -> list[dict]:
    if not TOOL_LOG.exists():
        return []
    lines = TOOL_LOG.read_text(encoding="utf-8").splitlines()[offset:]
    return [json.loads(line) for line in lines if line.strip()]


def docx_since(t: float) -> list[str]:
    if not OUT_DIR.exists():
        return []
    return sorted(str(p) for p in OUT_DIR.glob("*.docx") if p.stat().st_mtime >= t)


def docx_text(path: str) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


def field(body, *keys):
    """Safely read nested keys from a response that may not be a dict."""
    for k in keys:
        if not isinstance(body, dict):
            return None
        body = body.get(k)
    return body


class Run:
    def __init__(self):
        self.record = {
            "started": datetime.now(timezone.utc).isoformat(),
            "base": BASE,
            "model": MODEL,
            # "online" or "offline" -- set by stage0/offline_phase.ps1
            "label": os.environ.get("STAGE0_LABEL", "online"),
            "mode": os.environ.get("STAGE0_MODE", "full"),
            "steps": [],
        }

    def step(self, name: str, ok: bool, **detail) -> bool:
        self.record["steps"].append({"step": name, "ok": ok, **clip(detail)})
        secs = f"  ({detail['seconds']}s)" if "seconds" in detail else ""
        print(f"[{'PASS' if ok else 'FAIL'}] {name}{secs}")
        return ok

    def finish(self) -> int:
        self.record["finished"] = datetime.now(timezone.utc).isoformat()
        self.record["passed"] = all(s["ok"] for s in self.record["steps"])
        out = RESULTS / f"run_{self.record['label']}_{datetime.now():%Y%m%d-%H%M%S}.json"
        out.write_text(json.dumps(self.record, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"\n{'ALL STEPS PASSED' if self.record['passed'] else 'SOME STEPS FAILED'} -> {out}")
        return 0 if self.record["passed"] else 1


def upload(api: Api, path: Path, mime: str, slug: str):
    with path.open("rb") as f:
        code, body, secs = api.call(
            "POST", "/v1/document/upload",
            files={"file": (path.name, f, mime)},
            data={"addToWorkspaces": slug},
            timeout=UPLOAD_TIMEOUT,
        )
    docs = field(body, "documents") or []
    text = " ".join(d.get("pageContent", "") for d in docs if isinstance(d, dict))
    return code, body, secs, docs, text


def run_agent(api: Api, run: Run, slug: str, name: str, session: str, message: str,
              facts: dict | None = None) -> bool:
    """Start an agent through the API. With `facts`, the note's content must be right.

    Each test gets its own `session` so one test cannot reuse another's chat
    history -- run 2's Test B saw Test A's messages because they shared one.
    """
    offset, t0 = tool_log_len(), time.time()
    code, body, secs = api.call(
        "POST", f"/v1/workspace/{slug}/chat",
        json={"message": message, "mode": "chat", "sessionId": session},
        timeout=CHAT_TIMEOUT,
    )
    calls = [c for c in tool_calls_since(offset) if c.get("tool") == "create_word_document"]
    files = docx_since(t0)
    plumbing_ok = code == 200 and bool(calls) and bool(files)

    detail = {
        "status": code,
        "error": field(body, "error") if isinstance(body, dict) else body,
        "reply": field(body, "textResponse"),
        "tool_calls": calls,
        "new_docx": files,
        "plumbing_ok": plumbing_ok,
        "seconds": secs,
    }
    ok = plumbing_ok
    if facts is not None:
        text = docx_text(files[-1]) if files else ""
        checks = content_checks(text, facts)
        detail.update(docx_text=text, content_checks=checks)
        ok = plumbing_ok and all(checks.values())
    return run.step(name, ok, **detail)


def run_rag(api: Api, run: Run, slug: str, facts: dict) -> bool:
    """The trap question through plain document retrieval -- no agent."""
    cml, current, minimum = facts["breached"][0]
    question = (
        f"According to the inspection report for {facts['tag']}, which condition "
        f"monitoring location (CML) has a current thickness below the minimum "
        f"required, and what are its current and minimum thickness values?"
    )
    code, body, secs = api.call(
        "POST", f"/v1/workspace/{slug}/chat",
        json={"message": question, "mode": "query", "sessionId": "stage0-C"},
        timeout=CHAT_TIMEOUT,
    )
    reply = field(body, "textResponse") or ""
    checks = {
        f"names {cml}": cml in reply,
        f"gives current {current}": current in reply,
        f"gives minimum {minimum}": minimum in reply,
    }
    sources = [s.get("title") for s in (field(body, "sources") or []) if isinstance(s, dict)]
    return run.step(
        "test C (document Q&A, no agent): answers the trap question from the report",
        code == 200 and all(checks.values()),
        status=code,
        reply=reply,
        content_checks=checks,
        sources=sources,
        seconds=secs,
    )


def main() -> int:
    key = os.environ.get("ALLM_API_KEY") or KEY_FILE.read_text(encoding="utf-8").strip()
    api, run = Api(key), Run()
    facts = expected_facts()
    run.record["expected"] = facts

    code, body, secs = api.call("GET", "/v1/auth")
    if not run.step("API key accepted", code == 200, status=code, response=body, seconds=secs):
        return run.finish()

    code, body, secs = api.call(
        "POST", "/v1/workspace/new",
        json={"name": f"stage0-{datetime.now():%H%M%S}", "chatMode": "chat", "openAiTemp": 0.2},
    )
    slug = field(body, "workspace", "slug")
    if not run.step("workspace created", code == 200 and bool(slug), status=code, slug=slug, seconds=secs):
        return run.finish()

    code, body, secs = api.call(
        "POST", f"/v1/workspace/{slug}/update",
        json={"agentProvider": "ollama", "agentModel": MODEL},
    )
    run.step(
        "workspace agent model set",
        code == 200 and field(body, "workspace", "agentModel") == MODEL,
        status=code,
        agentProvider=field(body, "workspace", "agentProvider"),
        agentModel=field(body, "workspace", "agentModel"),
        seconds=secs,
    )

    code, body, secs, docs, text = upload(api, SCAN, "image/png", slug)
    run.step(
        "scanned report uploaded and embedded",
        code == 200 and bool(docs),
        status=code,
        error=field(body, "error") if isinstance(body, dict) else body,
        extracted_chars=len(text),
        extracted_preview=text[:600],
        seconds=secs,
    )
    scan_read = len(text) > 200 and facts["tag"] in text
    run.step(f"text extracted from the scan (finds '{facts['tag']}')", scan_read,
             extracted_chars=len(text))

    if run.record["mode"] == "upload-only":
        return run.finish()

    if not scan_read:
        code, body, secs, docs, text = upload(api, PDF, "application/pdf", slug)
        run.step(
            "fallback: born-digital PDF of the same report uploaded",
            code == 200 and facts["tag"] in text,
            status=code,
            extracted_chars=len(text),
            seconds=secs,
        )

    run_agent(
        api, run, slug,
        "test A (plumbing): agent calls the MCP tool; Word file lands on host",
        "stage0-A",
        "@agent Use the create_word_document tool to create a document titled "
        "'Stage 0 plumbing test' with the body 'Created by an AnythingLLM agent "
        "through the developer API.'",
    )

    run_agent(
        api, run, slug,
        "test B (realistic): agent writes an approval note with the right content",
        "stage0-B",
        "@agent Search this workspace's documents for the inspection report on "
        f"equipment {facts['tag']}. List each finding with its severity, and every "
        "thickness reading with its minimum required value. Then use the "
        "create_word_document tool to write a short approval note titled "
        f"'Approval Note {facts['tag']}' that summarises them.",
        facts=facts,
    )

    run_rag(api, run, slug, facts)

    return run.finish()


if __name__ == "__main__":
    raise SystemExit(main())
