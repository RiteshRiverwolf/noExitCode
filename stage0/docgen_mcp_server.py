"""Stage 0 test tool: a minimal MCP server that writes Word documents.

Runs on the host, not inside the AnythingLLM container, and serves MCP over
streamable HTTP so AnythingLLM in Docker can reach it:

    .venv\\Scripts\\python stage0\\docgen_mcp_server.py --port 8765

Endpoint: http://<host>:8765/mcp

Written against the MCP Python SDK v2, where v1's FastMCP became MCPServer and
transport settings (host, port, path, security) moved to run().

Every call is appended to results/stage0/tool_calls.jsonl, so a run can be
reconstructed afterwards without trusting AnythingLLM's own logs.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from mcp.server.mcpserver import MCPServer
from mcp.server.transport_security import TransportSecuritySettings

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "stage0" / "output"
LOG = ROOT / "results" / "stage0" / "tool_calls.jsonl"

# DNS-rebinding protection stays on. Bound to 127.0.0.1, the SDK by default only
# accepts the Host headers below minus host.docker.internal, which is how
# requests from a Docker Desktop container arrive -- so it is added explicitly.
ALLOWED_HOSTS = ["127.0.0.1:*", "localhost:*", "[::1]:*", "host.docker.internal:*"]
ALLOWED_ORIGINS = [f"http://{h}" for h in ALLOWED_HOSTS]


def _log(event: dict) -> None:
    LOG.parent.mkdir(parents=True, exist_ok=True)
    event["ts"] = datetime.now(timezone.utc).isoformat()
    with LOG.open("a", encoding="utf-8") as f:
        f.write(json.dumps(event) + "\n")


def build_server() -> MCPServer:
    mcp = MCPServer("sih-docgen")

    @mcp.tool()
    def create_word_document(title: str, body: str) -> str:
        """Create a Word (.docx) document on the workstation.

        Args:
            title: Document title, used as the heading and in the file name.
            body: Document text. Each line becomes a paragraph; lines starting
                with "- " become bullet points.

        Returns the saved file's name, path and SHA-256 hash.
        """
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        doc = Document()
        doc.add_heading(title.strip() or "Untitled", level=1)
        for raw in body.split("\n"):
            line = raw.rstrip()
            if not line:
                continue
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)

        safe = "".join(c if c.isalnum() or c in "-_" else "_" for c in title)[:60]
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        path = OUT_DIR / f"{safe or 'document'}_{stamp}.docx"
        doc.save(path)
        digest = hashlib.sha256(path.read_bytes()).hexdigest()

        _log({
            "tool": "create_word_document",
            "title": title,
            "body_chars": len(body),
            "path": str(path),
            "sha256": digest,
        })
        return f"Saved {path.name} (sha256 {digest[:16]}...) at {path}"

    return mcp


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--host", default="127.0.0.1")
    ap.add_argument("--port", type=int, default=8765)
    args = ap.parse_args()

    security = TransportSecuritySettings(
        enable_dns_rebinding_protection=True,
        allowed_hosts=ALLOWED_HOSTS,
        allowed_origins=ALLOWED_ORIGINS,
    )
    _log({"event": "server_start", "host": args.host, "port": args.port,
          "allowed_hosts": ALLOWED_HOSTS})
    build_server().run(
        transport="streamable-http",
        host=args.host,
        port=args.port,
        streamable_http_path="/mcp",
        transport_security=security,
    )


if __name__ == "__main__":
    main()
